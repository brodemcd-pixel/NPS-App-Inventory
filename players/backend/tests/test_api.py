"""API integration tests against real Postgres + pgvector (see conftest fixture)."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

pytestmark = pytest.mark.anyio


def _parse_sse(body: str) -> list[tuple[str, str]]:
    events = []
    for block in body.strip().split("\n\n"):
        event, data = None, ""
        for line in block.split("\n"):
            if line.startswith("event: "):
                event = line[len("event: "):]
            elif line.startswith("data: "):
                data = line[len("data: "):]
        if event:
            events.append((event, data))
    return events


class TestPlayersList:
    async def test_filters_and_sort(self, client):
        r = await client.get("/api/players", params={"position_group": "WR",
                                                     "sort": "forty"})
        assert r.status_code == 200
        body = r.json()
        assert body["total"] == 4
        fortys = [p["forty"] for p in body["items"]]
        assert fortys == sorted(fortys)
        assert body["items"][0]["name"] == "Ace Alpha"
        assert body["items"][0]["scheme_archetype"] == "Air Raid"
        assert "slot" in body["items"][0]["roles"]

    async def test_class_round_and_measurable_filters(self, client):
        r = await client.get("/api/players", params=[("draft_class", 2024),
                                                     ("round", 1)])
        names = {p["name"] for p in r.json()["items"]}
        assert names == {"Ace Alpha", "Ed Echo"}

        r = await client.get("/api/players", params={"forty_max": 4.4})
        assert {p["name"] for p in r.json()["items"]} == {"Ace Alpha", "Ed Echo"}

    async def test_trait_and_flag_filters(self, client):
        r = await client.get("/api/players", params={"trait": "motor",
                                                     "trait_min": 0.7})
        names = {p["name"] for p in r.json()["items"]}
        assert names == {"Ace Alpha", "Ike India", "Jay Juliett"}

        r = await client.get("/api/players", params={"red_flag": True})
        assert [p["name"] for p in r.json()["items"]] == ["Bo Bravo"]

    async def test_scheme_and_role_filters(self, client):
        r = await client.get("/api/players", params={"scheme": "Air Raid"})
        assert {p["name"] for p in r.json()["items"]} == {"Ace Alpha", "Bo Bravo"}
        r = await client.get("/api/players", params={"role": "press-man corner"})
        assert [p["name"] for p in r.json()["items"]] == ["Ed Echo"]

    async def test_pagination(self, client):
        r = await client.get("/api/players", params={"page_size": 3, "page": 2,
                                                     "sort": "name"})
        body = r.json()
        assert body["total"] == 10 and len(body["items"]) == 3
        assert body["page"] == 2


class TestPlayerDetail:
    async def test_detail_shape(self, client):
        r = await client.get("/api/players/1")
        assert r.status_code == 200
        d = r.json()
        assert d["name"] == "Ace Alpha"
        for key in ("measurables", "report", "flags", "scheme_context", "roles",
                    "mental_profile", "production", "trajectory", "nfl_outcomes",
                    "mock_consensus", "comments", "provenance", "umap"):
            assert key in d, key
        assert d["scheme_context"][-1]["scheme_archetype"] == "Air Raid"
        traits = {t["trait"] for t in d["mental_profile"]["core"]}
        assert {"anticipation", "motor", "football_iq"} <= traits
        assert d["trajectory"]["trend"] in ("ascending", "flat", "declining")
        assert len(d["mock_consensus"]) == 4

    async def test_404(self, client):
        assert (await client.get("/api/players/999")).status_code == 404


class TestSimilarity:
    async def test_same_position_group_default(self, client):
        r = await client.get("/api/players/1/similar", params={"limit": 5})
        body = r.json()
        assert body["weights_used"] == {"scouting": 0.5, "scheme": 0.25, "mental": 0.25}
        assert all(i["player"]["position_group"] == "WR" for i in body["items"])
        # Bo Bravo: same school/scheme/roles overlap + shared report vocabulary -> top comp
        assert body["items"][0]["player"]["name"] == "Bo Bravo"
        overall = [i["overall"] for i in body["items"]]
        assert overall == sorted(overall, reverse=True)

    async def test_weights_change_ranking_inputs(self, client):
        default = (await client.get("/api/players/1/similar")).json()
        scouting_only = (
            await client.get("/api/players/1/similar", params={"w_scouting": 1,
                                                               "w_scheme": 0,
                                                               "w_mental": 0})
        ).json()
        assert scouting_only["weights_used"]["scouting"] == 1.0
        a = {i["player"]["id"]: i["overall"] for i in default["items"]}
        b = {i["player"]["id"]: i["overall"] for i in scouting_only["items"]}
        assert any(abs(a[k] - b[k]) > 1e-6 for k in set(a) & set(b))

    async def test_mental_axis_null_dropped(self, client):
        # Hal Hotel (id 8) has no mental rows; axis must be null for him.
        r = await client.get("/api/players/7/similar", params={"limit": 5})
        items = r.json()["items"]
        hal = next(i for i in items if i["player"]["name"] == "Hal Hotel")
        assert hal["axes"]["mental"] is None
        assert hal["overall"] > 0  # blended from remaining axes

    async def test_compare(self, client):
        r = await client.get("/api/compare", params={"a": 1, "b": 2})
        body = r.json()
        assert body["a"]["name"] == "Ace Alpha" and body["b"]["name"] == "Bo Bravo"
        assert 0 <= body["similarity"]["overall"] <= 1
        assert set(body["similarity"]["axes"]) == {"scouting", "scheme", "mental"}

    async def test_team_fit(self, client):
        r = await client.get("/api/players/1/team_fit", params={"limit": 2})
        items = r.json()["items"]
        assert items[0]["team"] == "Las Vegas Raiders"  # Air Raid + Leach tree match
        assert items[0]["fit"] > items[-1]["fit"] or len(items) == 1


class TestWatchlists:
    async def test_crud_flow(self, client):
        headers = {"X-User-Id": "scout-7"}
        created = (await client.post("/api/watchlists", json={"name": "Sleepers"},
                                     headers=headers)).json()
        wid = created["id"]
        r = await client.post(f"/api/watchlists/{wid}/players",
                              json={"player_id": 1, "note": "deep threat"},
                              headers=headers)
        assert r.status_code == 201
        lists = (await client.get("/api/watchlists", headers=headers)).json()["items"]
        assert lists[0]["players"][0]["name"] == "Ace Alpha"
        assert lists[0]["players"][0]["note"] == "deep threat"
        # other users see nothing
        other = (await client.get("/api/watchlists",
                                  headers={"X-User-Id": "someone-else"})).json()
        assert other["items"] == []
        r = await client.delete(f"/api/watchlists/{wid}/players/1", headers=headers)
        assert r.status_code == 204
        r = await client.delete(f"/api/watchlists/{wid}", headers=headers)
        assert r.status_code == 204
        assert (await client.get("/api/watchlists", headers=headers)).json()["items"] == []


class TestComments:
    async def test_comment_creates_scout_tag(self, client):
        r = await client.post(
            "/api/comments",
            json={"player_id": 1, "body": "Great interview",
                  "traits": [{"trait": "leadership", "score": 0.9},
                             {"trait": "bogus", "score": 0.5}]},
            headers={"X-User-Id": "scout-9"},
        )
        assert r.status_code == 201
        assert r.json()["traits"] == [{"trait": "leadership", "score": 0.9}]
        detail = (await client.get("/api/players/1")).json()
        leadership = next(t for t in detail["mental_profile"]["core"]
                          if t["trait"] == "leadership")
        assert "scout_tag" in leadership["sources"]
        assert detail["comments"][0]["body"] == "Great interview"


class TestExport:
    async def test_xlsx(self, client):
        r = await client.get("/api/export/players.xlsx",
                             params={"position_group": "WR"})
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["content-type"]
        import io
        ws = load_workbook(io.BytesIO(r.content)).active
        assert ws.max_row == 5  # header + 4 WRs

    async def test_pdf_501_without_weasyprint(self, client):
        r = await client.get("/api/export/players/1.pdf")
        assert r.status_code == 501
        assert "WeasyPrint" in r.json()["detail"]


class TestChatOffline:
    async def test_sse_sequence_and_sources(self, client):
        r = await client.post("/api/chat",
                              json={"message": "3 fastest 40 times among receivers"})
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("text/event-stream")
        events = _parse_sse(r.text)
        kinds = [e for e, _ in events]
        assert kinds[0] == "tool"
        assert "text" in kinds and kinds[-2:] == ["sources", "done"]
        import json
        sources = json.loads(dict(events)["sources"])["sources"]
        assert sources[0]["operation"] == "filtered_list"
        assert sources[0]["count"] == 3
        text_payload = "".join(
            json.loads(d)["delta"] for e, d in events if e == "text"
        )
        assert "Ace Alpha" in text_payload and "/players/1" in text_payload

    async def test_capabilities_fallback(self, client):
        r = await client.post("/api/chat", json={"message": "zzz qqq unknowable"})
        events = _parse_sse(r.text)
        assert [e for e, _ in events][-2:] == ["sources", "done"]


class TestDashboards:
    async def test_accuracy_shape(self, client):
        r = await client.get("/api/accuracy")
        body = r.json()
        assert body["classes"] == [2022, 2023, 2024]
        names = [m["name"] for m in body["models"]]
        assert names == ["FSM v0.2", "FSM v1.0"]
        assert all("spearman" in m and "mae" in m and "n" in m for m in body["models"])

    async def test_analytics_summary(self, client):
        await client.get("/api/players/1")  # generates a player_view event
        await client.post("/api/events", json={"event_type": "page_view",
                                               "payload": {"path": "/"}})
        r = await client.get("/api/analytics/summary")
        body = r.json()
        assert body["page_views"] >= 1
        assert {"player_id", "name", "count"} <= set(body["most_discussed"][0])

    async def test_umap(self, client):
        items = (await client.get("/api/umap")).json()["items"]
        assert len(items) == 10
        assert {"player_id", "x", "y", "position_group"} <= set(items[0])


class TestAuth:
    async def test_slack_token_roundtrip(self, client):
        from backend.app.services.slack_link import sign_slack_user

        token = sign_slack_user("U123", "Scout Jane")
        r = await client.post("/api/auth/slack/exchange", json={"token": token})
        assert r.status_code == 200
        assert r.json() == {"user_id": "U123", "name": "Scout Jane"}
        r = await client.post("/api/auth/slack/exchange", json={"token": "garbage"})
        assert r.status_code == 401

    async def test_meta_filters(self, client):
        body = (await client.get("/api/meta/filters")).json()
        assert 2024 in body["classes"]
        assert "Air Raid" in body["offense_schemes"]
        assert "football_iq" in body["traits"]
